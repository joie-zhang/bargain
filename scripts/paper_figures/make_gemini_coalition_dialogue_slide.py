#!/usr/bin/env python3
"""Create a one-slide dialogue deck for the Gemini coalition case studies."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "analysis" / "figure_recreation_20260816"
PNG_PATH = OUTPUT_DIR / "gemini_coalition_dialogue_slide.png"
PPTX_PATH = OUTPUT_DIR / "gemini_coalition_dialogue_slide.pptx"

WIDTH = 2400
HEIGHT = 855
PANEL_WIDTHS = (665, 609, 1126)

COLORS = {
    "navy": "#142843",
    "navy_2": "#203A5D",
    "background": "#F3F6FA",
    "card": "#FFFFFF",
    "border": "#CBD5E1",
    "text": "#172033",
    "muted": "#5C687A",
    "blue": "#2F6FED",
    "blue_dark": "#2056BF",
    "blue_bg": "#EDF4FF",
    "red": "#D64545",
    "red_dark": "#A92F36",
    "red_bg": "#FFF0F0",
    "green": "#16805B",
    "green_bg": "#E9F7EF",
    "white_blue": "#D5E2F3",
    "purple": "#7C3AED",
    "teal": "#0F8F83",
    "magenta": "#C24183",
    "gold": "#B7791F",
    "cyan": "#087EA4",
    "orange": "#E26A2C",
    "orange_bg": "#FFF2E8",
}

CASES = [
    {
        "config": "RUN 107",
        "threshold": "3 of 4 votes needed\nfor supermajority",
        "coalition": [1, 2, 3],
        "excluded": [4],
        "planner": 1,
        "agent_colors": {
            1: "blue",
            2: "purple",
            3: "teal",
            4: "red",
        },
        "plan": (
            "“A supermajority only requires 3 votes. I can secure a guaranteed "
            "supermajority by forming an alliance with Agent_2 and Agent_3 ... "
            "leaving Agent_4 with just the Clock.”"
        ),
        "responses": [
            (
                4,
                "“This proposal only gives me the Clock, resulting in a very low "
                "utility of 20.”",
            )
        ],
        "result": "PASSED 3–1",
        "payoffs": "AGENT PAYOFFS  [83, 72, 81, 20]",
        "red_payoff_indices": (3,),
    },
    {
        "config": "RUN 108",
        "threshold": "3 of 4 votes needed\nfor supermajority",
        "coalition": [1, 2, 3],
        "excluded": [4],
        "planner": 2,
        "agent_colors": {
            1: "purple",
            2: "blue",
            3: "teal",
            4: "red",
        },
        "plan": (
            "“My best path forward is to form a 3-agent voting bloc with Agent_1 "
            "and Agent_3, effectively freezing out Agent_4.”"
        ),
        "responses": [
            (
                4,
                "“This proposal gives me absolutely nothing, resulting in zero utility.”",
            )
        ],
        "result": "PASSED 3–1",
        "payoffs": "AGENT PAYOFFS  [49, 73, 55, 0]",
        "red_payoff_indices": (3,),
    },
    {
        "config": "RUN 117",
        "threshold": "6 of 8 votes needed\nfor supermajority",
        "coalition": [1, 3, 4, 5, 6, 8],
        "excluded": [2, 7],
        "planner": 8,
        "agent_colors": {
            1: "purple",
            2: "red",
            3: "teal",
            4: "magenta",
            5: "gold",
            6: "cyan",
            7: "orange",
            8: "blue",
        },
        "plan": (
            "“I present the Perfect Alliance of 6. By excluding them, we have exactly "
            "6 votes where NO ONE in the Core 5 has to compromise.”"
        ),
        "responses": [
            (
                2,
                "“This proposal gives me absolutely zero items, resulting in a utility "
                "of 0. I cannot accept it.”",
            ),
            (7, "“This proposal gives me zero items, resulting in a utility of 0.”"),
        ],
        "result": "PASSED 6–2",
        "payoffs": "AGENT PAYOFFS  [51, 0, 69, 59, 50, 64, 0, 52]",
        "red_payoff_indices": (1, 6),
    },
]


FONT_REGULAR = "/usr/share/fonts/adobe-source-sans-pro-fonts/SourceSans3-Regular.otf"
FONT_BOLD = "/usr/share/fonts/adobe-source-sans-pro-fonts/SourceSans3-Bold.otf"


def font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size=size)


def box(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    *,
    fill: str | None,
    outline: str | None = None,
    width: int = 1,
) -> None:
    draw.rectangle(bounds, fill=fill, outline=outline, width=width)


def wrap_lines(
    draw: ImageDraw.ImageDraw,
    text: str,
    text_font: ImageFont.FreeTypeFont,
    max_width: int,
) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), trial, font=text_font)[2] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    *,
    text_font: ImageFont.FreeTypeFont,
    fill: str,
    max_width: int,
    line_gap: int = 5,
) -> int:
    x, y = xy
    lines = wrap_lines(draw, text, text_font, max_width)
    line_height = text_font.getbbox("Ag")[3] - text_font.getbbox("Ag")[1]
    for line in lines:
        draw.text((x, y), line, font=text_font, fill=fill)
        y += line_height + line_gap
    return y


def draw_payoffs(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    *,
    text_font: ImageFont.FreeTypeFont,
    red_indices: tuple[int, ...],
) -> None:
    prefix, separator, values_text = text.partition("[")
    if not separator or not values_text.endswith("]"):
        raise ValueError(f"Unexpected payoff text: {text}")

    x, y = xy

    def draw_run(run: str, color: str) -> None:
        nonlocal x
        draw.text((x, y), run, font=text_font, fill=color)
        x += round(draw.textlength(run, font=text_font))

    draw_run(f"{prefix}[", COLORS["text"])
    values = [value.strip() for value in values_text[:-1].split(",")]
    for index, value in enumerate(values):
        if index:
            draw_run(", ", COLORS["text"])
        color = COLORS["red"] if index in red_indices else COLORS["text"]
        draw_run(value, color)
    draw_run("]", COLORS["text"])


def draw_robot_avatar(
    draw: ImageDraw.ImageDraw,
    center: tuple[int, int],
    *,
    color: str,
    radius: int,
) -> None:
    cx, cy = center
    draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=color)

    head_w = round(radius * 1.10)
    head_h = round(radius * 0.82)
    left = cx - head_w // 2
    top = cy - head_h // 2 + 3
    draw.rectangle((left, top, left + head_w, top + head_h), fill="white")
    antenna_top = cy - radius + max(5, radius // 6)
    draw.line((cx, top, cx, antenna_top), fill="white", width=max(2, radius // 8))
    dot = max(2, radius // 9)
    draw.ellipse((cx - dot, antenna_top - dot, cx + dot, antenna_top + dot), fill="white")
    eye_r = max(2, radius // 9)
    eye_y = top + round(head_h * 0.43)
    for eye_x in (cx - round(head_w * 0.22), cx + round(head_w * 0.22)):
        draw.ellipse(
            (eye_x - eye_r, eye_y - eye_r, eye_x + eye_r, eye_y + eye_r),
            fill=color,
        )
    mouth_y = top + round(head_h * 0.70)
    draw.line(
        (cx - round(head_w * 0.20), mouth_y, cx + round(head_w * 0.20), mouth_y),
        fill=color,
        width=max(2, radius // 9),
    )


def draw_agent_chip(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    agent: int,
    *,
    color: str,
) -> int:
    draw_robot_avatar(draw, (x + 25, y + 25), color=color, radius=25)
    draw.text((x + 58, y + 2), f"A{agent}", font=font(34, bold=True), fill=COLORS["text"])
    return x + 105


def draw_dialogue_bubble(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    *,
    agent: int,
    quote: str,
    speaker_role: str,
    kind: str,
    speaker_color: str,
    quote_size: int,
    label_size: int = 27,
) -> None:
    left, top, right, bottom = bounds
    if kind == "plan":
        color = COLORS["blue"]
        fill = COLORS["blue_bg"]
        label_color = COLORS["blue_dark"]
    elif speaker_color == COLORS["orange"]:
        color = speaker_color
        fill = COLORS["orange_bg"]
        label_color = "#A94518"
    else:
        color = speaker_color
        fill = COLORS["red_bg"]
        label_color = COLORS["red_dark"]

    bubble_left = left + 82
    draw.polygon(
        [(bubble_left, top + 42), (bubble_left - 24, top + 60), (bubble_left, top + 78)],
        fill=fill,
    )
    box(draw, (bubble_left, top, right, bottom), fill=fill)
    draw_robot_avatar(draw, (left + 34, top + 56), color=color, radius=34)
    draw.text(
        (bubble_left + 22, top + 13),
        f"AGENT {agent}  ·  {speaker_role}",
        font=font(label_size, bold=True),
        fill=label_color,
    )
    draw_wrapped(
        draw,
        quote,
        (bubble_left + 22, top + 56),
        text_font=font(quote_size),
        fill=COLORS["text"],
        max_width=right - bubble_left - 42,
        line_gap=4,
    )


def dialogue_box_height(
    draw: ImageDraw.ImageDraw,
    quote: str,
    *,
    text_width: int,
    quote_size: int,
) -> int:
    quote_font = font(quote_size)
    line_height = quote_font.getbbox("Ag")[3] - quote_font.getbbox("Ag")[1]
    line_count = len(wrap_lines(draw, quote, quote_font, text_width))
    return 56 + line_count * (line_height + 4) + 12


def build_slide_image() -> Image.Image:
    image = Image.new("RGB", (WIDTH, HEIGHT), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    card_top = 0
    card_bottom = HEIGHT
    dialogue_top = 279
    minimum_bubble_gap = 18
    quote_size = 32

    panel_layouts = []
    for case, panel_width in zip(CASES, PANEL_WIDTHS, strict=True):
        text_width = panel_width - 176
        quotes = [case["plan"], *(quote for _, quote in case["responses"])]
        bubble_heights = [
            dialogue_box_height(
                draw,
                quote,
                text_width=text_width,
                quote_size=quote_size,
            )
            for quote in quotes
        ]
        minimum_group_height = sum(bubble_heights) + minimum_bubble_gap * (
            len(bubble_heights) - 1
        )
        panel_layouts.append((bubble_heights, minimum_group_height))

    dialogue_group_height = max(group_height for _, group_height in panel_layouts)
    result_top = dialogue_top + dialogue_group_height + 22
    result_bottom = HEIGHT - 20

    left = 0
    for index, (case, panel_width) in enumerate(zip(CASES, PANEL_WIDTHS, strict=True)):
        right = left + panel_width
        box(draw, (left, card_top, right, card_bottom), fill=COLORS["card"])

        inner_left = left + 26
        inner_right = right - 26

        draw.text((inner_left, -3), case["config"], font=font(48, bold=True), fill=COLORS["navy"])
        vote_text = case["threshold"]
        vote_font = font(28, bold=True)
        vote_bbox = draw.multiline_textbbox((0, 0), vote_text, font=vote_font, spacing=0, align="center")
        vote_text_width = vote_bbox[2] - vote_bbox[0]
        draw.multiline_text(
            (inner_right - vote_text_width, 3),
            vote_text,
            font=vote_font,
            fill=COLORS["muted"],
            spacing=0,
            align="center",
        )

        draw.text((inner_left, 79), "WINNING BLOC", font=font(27, bold=True), fill=COLORS["muted"])
        chip_x = inner_left
        for agent in case["coalition"]:
            agent_color = COLORS[case["agent_colors"][agent]]
            chip_x = draw_agent_chip(draw, chip_x, 111, agent, color=agent_color)

        draw.text((inner_left, 171), "LEFT OUT", font=font(27, bold=True), fill=COLORS["muted"])
        chip_x = inner_left
        for agent in case["excluded"]:
            agent_color = COLORS[case["agent_colors"][agent]]
            chip_x = draw_agent_chip(draw, chip_x, 203, agent, color=agent_color)

        bubbles = [
            (
                case["planner"],
                case["plan"],
                "COALITION PLAN",
                "plan",
                COLORS[case["agent_colors"][case["planner"]]],
            ),
            *(
                (
                    response_agent,
                    response_quote,
                    "REJECTS",
                    "reject",
                    COLORS[case["agent_colors"][response_agent]],
                )
                for response_agent, response_quote in case["responses"]
            ),
        ]
        bubble_heights = panel_layouts[index][0]
        total_gap = dialogue_group_height - sum(bubble_heights)
        gap_count = len(bubble_heights) - 1
        bubble_gap, extra_gap_pixels = divmod(total_gap, gap_count)
        bubble_top = dialogue_top
        for bubble_index, (bubble, bubble_height) in enumerate(
            zip(bubbles, bubble_heights, strict=True)
        ):
            agent, quote, speaker_role, kind, speaker_color = bubble
            draw_dialogue_bubble(
                draw,
                (inner_left, bubble_top, inner_right, bubble_top + bubble_height),
                agent=agent,
                quote=quote,
                speaker_role=speaker_role,
                kind=kind,
                speaker_color=speaker_color,
                quote_size=quote_size,
            )
            bubble_top += bubble_height
            if bubble_index < gap_count:
                bubble_top += bubble_gap + (bubble_index < extra_gap_pixels)

        if bubble_top != dialogue_top + dialogue_group_height:
            raise RuntimeError("Dialogue groups must have equal rendered heights")

        box(
            draw,
            (inner_left, result_top, inner_right, result_bottom),
            fill=COLORS["green_bg"],
        )
        result_size = 31
        payoff_size = 36
        draw.text(
            (inner_left + 22, result_top + 16),
            case["result"],
            font=font(result_size, bold=True),
            fill=COLORS["green"],
        )
        draw_payoffs(
            draw,
            case["payoffs"],
            (inner_left + 22, result_top + 65),
            text_font=font(payoff_size, bold=True),
            red_indices=case["red_payoff_indices"],
        )

        left = right

    return image


def save_pptx(image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(13.333333 * HEIGHT / WIDTH)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(PPTX_PATH)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    image = build_slide_image()
    image.save(PNG_PATH, dpi=(180, 180))
    save_pptx(PNG_PATH)
    print(PNG_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
