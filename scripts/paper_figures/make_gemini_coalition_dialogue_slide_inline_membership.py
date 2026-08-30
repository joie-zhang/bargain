#!/usr/bin/env python3
"""Create the Gemini coalition slide with membership in the plan labels."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

import make_gemini_coalition_dialogue_slide as base


OUTPUT_DIR = base.OUTPUT_DIR
PNG_PATH = OUTPUT_DIR / "gemini_coalition_dialogue_slide_inline_membership.png"
PPTX_PATH = OUTPUT_DIR / "gemini_coalition_dialogue_slide_inline_membership.pptx"

PANEL_WIDTHS = (641, 585, 800)
WIDTH = sum(PANEL_WIDTHS)
HEIGHT = 665
HORIZONTAL_INSET = 14
PLAN_ROLES = (
    "COALITION PLAN FOR A1-A3",
    "COALITION PLAN FOR A1-A3",
    "COALITION PLAN EXCLUDING A2 AND A7",
)


def build_slide_image() -> Image.Image:
    image = Image.new("RGB", (WIDTH, HEIGHT), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    dialogue_top = 89
    minimum_bubble_gap = 18
    quote_size = 32

    panel_layouts = []
    for case, panel_width in zip(base.CASES, PANEL_WIDTHS, strict=True):
        text_width = panel_width - 2 * HORIZONTAL_INSET - 124
        quotes = [case["plan"], *(quote for _, quote in case["responses"])]
        bubble_heights = [
            base.dialogue_box_height(
                draw,
                quote,
                text_width=text_width,
                quote_size=quote_size,
            )
            for quote in quotes
        ]
        minimum_group_height = sum(bubble_heights) + minimum_bubble_gap * (
            len(bubble_heights) - 1
        )
        panel_layouts.append((bubble_heights, minimum_group_height))

    dialogue_group_height = max(
        group_height for _, group_height in panel_layouts[:2]
    )
    result_top = dialogue_top + dialogue_group_height + 22
    result_bottom = HEIGHT - 20

    left = 0
    for index, (case, panel_width, plan_role) in enumerate(
        zip(base.CASES, PANEL_WIDTHS, PLAN_ROLES, strict=True)
    ):
        right = left + panel_width
        inner_left = left + HORIZONTAL_INSET
        inner_right = right - HORIZONTAL_INSET

        base.box(draw, (left, 0, right, HEIGHT), fill=base.COLORS["card"])
        draw.text(
            (inner_left, -3),
            case["config"],
            font=base.font(48, bold=True),
            fill=base.COLORS["navy"],
        )

        vote_font = base.font(28, bold=True)
        vote_box = draw.multiline_textbbox(
            (0, 0),
            case["threshold"],
            font=vote_font,
            spacing=0,
            align="center",
        )
        vote_width = vote_box[2] - vote_box[0]
        draw.multiline_text(
            (inner_right - vote_width, 3),
            case["threshold"],
            font=vote_font,
            fill=base.COLORS["muted"],
            spacing=0,
            align="center",
        )

        bubbles = [
            (
                case["planner"],
                case["plan"],
                plan_role,
                "plan",
                base.COLORS[case["agent_colors"][case["planner"]]],
            ),
            *(
                (
                    response_agent,
                    response_quote,
                    "REJECTS",
                    "reject",
                    base.COLORS[case["agent_colors"][response_agent]],
                )
                for response_agent, response_quote in case["responses"]
            ),
        ]
        bubble_heights = panel_layouts[index][0]
        if index == 2:
            agent, quote, speaker_role, kind, speaker_color = bubbles[0]
            base.draw_dialogue_bubble(
                draw,
                (
                    inner_left,
                    dialogue_top,
                    inner_right,
                    dialogue_top + bubble_heights[0],
                ),
                agent=agent,
                quote=quote,
                speaker_role=speaker_role,
                kind=kind,
                speaker_color=speaker_color,
                quote_size=quote_size,
                label_size=26,
            )

            pair_gap = 16
            pair_total_width = inner_right - inner_left - pair_gap
            right_pair_width = 360
            left_pair_width = pair_total_width - right_pair_width
            pair_bounds = (
                (inner_left, inner_left + left_pair_width),
                (inner_left + left_pair_width + pair_gap, inner_right),
            )
            pair_heights = [
                base.dialogue_box_height(
                    draw,
                    bubble[1],
                    text_width=pair_right - pair_left - 124,
                    quote_size=quote_size,
                )
                for bubble, (pair_left, pair_right) in zip(
                    bubbles[1:], pair_bounds, strict=True
                )
            ]
            pair_height = max(pair_heights)
            vertical_gap = (
                result_top - dialogue_top - bubble_heights[0] - pair_height
            ) // 2
            pair_top = dialogue_top + bubble_heights[0] + vertical_gap

            for bubble, (pair_left, pair_right) in zip(
                bubbles[1:], pair_bounds, strict=True
            ):
                agent, quote, speaker_role, kind, speaker_color = bubble
                base.draw_dialogue_bubble(
                    draw,
                    (pair_left, pair_top, pair_right, pair_top + pair_height),
                    agent=agent,
                    quote=quote,
                    speaker_role=speaker_role,
                    kind=kind,
                    speaker_color=speaker_color,
                    quote_size=quote_size,
                    label_size=26,
                )
        else:
            total_gap = dialogue_group_height - sum(bubble_heights)
            gap_count = len(bubble_heights) - 1
            bubble_gap, extra_gap_pixels = divmod(total_gap, gap_count)
            bubble_top = dialogue_top

            for bubble_index, (bubble, bubble_height) in enumerate(
                zip(bubbles, bubble_heights, strict=True)
            ):
                agent, quote, speaker_role, kind, speaker_color = bubble
                base.draw_dialogue_bubble(
                    draw,
                    (inner_left, bubble_top, inner_right, bubble_top + bubble_height),
                    agent=agent,
                    quote=quote,
                    speaker_role=speaker_role,
                    kind=kind,
                    speaker_color=speaker_color,
                    quote_size=quote_size,
                    label_size=26,
                )
                bubble_top += bubble_height
                if bubble_index < gap_count:
                    bubble_top += bubble_gap + (bubble_index < extra_gap_pixels)

            if bubble_top != dialogue_top + dialogue_group_height:
                raise RuntimeError("Dialogue groups must have equal rendered heights")

        base.box(
            draw,
            (inner_left, result_top, inner_right, result_bottom),
            fill=base.COLORS["green_bg"],
        )
        draw.text(
            (inner_left + 22, result_top + 16),
            case["result"],
            font=base.font(31, bold=True),
            fill=base.COLORS["green"],
        )
        base.draw_payoffs(
            draw,
            case["payoffs"],
            (inner_left + 22, result_top + 65),
            text_font=base.font(36, bold=True),
            red_indices=case["red_payoff_indices"],
        )

        left = right

    return image


def save_pptx(image_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(13.333333 * HEIGHT / WIDTH)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(PPTX_PATH)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    image = build_slide_image()
    image.save(PNG_PATH, dpi=(180, 180))
    save_pptx(PNG_PATH)
    print(PNG_PATH)
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
